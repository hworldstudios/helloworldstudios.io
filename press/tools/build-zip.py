#!/usr/bin/env python3

import html
import io
import json
import re
import zipfile
from pathlib import Path

from PIL import Image

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image as PdfImage,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    HAVE_PDF = True
except ImportError:
    HAVE_PDF = False

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


press = Path(__file__).resolve().parent.parent
kit = press / "press-kit"
source = press / "presskit.md"
archive = press / "press-kit.zip"
manifest = press / "manifest.json"
logo = kit / "logos" / "beltfed-logo-color.png"

SKIP_NAMES = {".ds_store", "thumbs.db", ".gitkeep"}
ASSET_FOLDERS = ("screenshots", "logos", "key-art", "team")
CLIPS_DIR = press / "gifs"
SITE_CLIPS = "https://helloworldstudios.io/press/gifs/"

INK = "#1A1A1A"
MUTED = "#6B6B6B"
ACCENT = "#C2410C"
RULE = "#D8D4D0"
TODO_INK = "#C2185B"
TODO_BG = "#FDE9F1"

TOKEN_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)|\*\*([^*]+)\*\*|`([^`]+)`")
MARKER_RE = re.compile(r"<!--\s*(gallery|youtube|clips)\s*:\s*([^\s>]+)\s*-->")
COMMENT_RE = re.compile(r"<!--.*?-->", re.S)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
PAIR_RE = re.compile(r"^\*\*([^*]+):\*\*\s*(.*)$")
BULLET_RE = re.compile(r"^[-*]\s+")
BLOCK_END_RE = re.compile(r"^(#{1,6}\s|[-*]\s|>|@@M\d+@@|-{3,}$)")


def runs(text):
    out, pos = [], 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            out.append({"text": text[pos:match.start()]})
        if match.group(1) is not None:
            out.append({"text": match.group(1), "href": match.group(2)})
        elif match.group(3) is not None:
            out.append({"text": match.group(3), "bold": True})
        else:
            out.append({"text": match.group(4), "code": True})
        pos = match.end()
    if pos < len(text):
        out.append({"text": text[pos:]})
    return out


def plain(spans):
    return "".join(s["text"] for s in spans)


def is_soon(spans):
    return plain(spans).strip().lower() == "coming soon"


def collect_clips():
    if not CLIPS_DIR.is_dir():
        return []
    clips = []
    for webm in sorted(CLIPS_DIR.glob("*.webm")):
        gif = webm.with_suffix(".gif")
        files = [("WebM", webm.name, webm.stat().st_size)]
        if gif.exists():
            files.append(("GIF", gif.name, gif.stat().st_size))
        clips.append({"stem": webm.stem, "files": files})
    return clips


def format_bytes(size):
    mb = size / 1048576
    return f"{mb:.1f} MB" if mb >= 1 else f"{round(size / 1024)} KB"


def parse(markdown, listing):
    markers = []

    def stash(match):
        markers.append((match.group(1), match.group(2)))
        return f"\n@@M{len(markers) - 1}@@\n"

    text = COMMENT_RE.sub("", MARKER_RE.sub(stash, markdown.replace("\r\n", "\n")))
    lines = text.split("\n")

    title, lead, blocks = "", "", []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        marker = re.fullmatch(r"@@M(\d+)@@", line)
        if marker:
            kind, value = markers[int(marker.group(1))]
            if kind == "gallery":
                files = listing.get(value, [])
                blocks.append({
                    "type": "files",
                    "folder": value,
                    "files": [(f.name, f.stat().st_size) for f in files],
                })
            elif kind == "clips":
                blocks.append({"type": "clips", "clips": collect_clips()})
            i += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level, body = len(heading.group(1)), heading.group(2).strip()
            if level == 1:
                title = body
            else:
                blocks.append({"type": "heading", "level": level, "spans": runs(body)})
            i += 1
            continue

        if re.fullmatch(r"-{3,}|\*{3,}", line):
            blocks.append({"type": "rule"})
            i += 1
            continue

        if BULLET_RE.match(line):
            items = []
            while i < len(lines) and BULLET_RE.match(lines[i].strip()):
                items.append(BULLET_RE.sub("", lines[i].strip()))
                i += 1
            pairs = [PAIR_RE.match(item) for item in items]
            if all(pairs):
                blocks.append({
                    "type": "facts",
                    "rows": [(m.group(1), runs(m.group(2))) for m in pairs],
                })
            else:
                blocks.append({"type": "list", "items": [runs(x) for x in items]})
            continue

        if line.startswith(">"):
            quoted = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quoted.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            attribution = ""
            if quoted and quoted[-1].startswith("-"):
                attribution = quoted.pop().lstrip("- ").strip()
            blocks.append({
                "type": "quote",
                "spans": runs(" ".join(x for x in quoted if x)),
                "attribution": attribution,
            })
            continue

        block = []
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate or BLOCK_END_RE.match(candidate):
                break
            block.append(candidate)
            i += 1

        if block[0].startswith("TODO:"):
            items = []
            for entry in block:
                if not entry.startswith("TODO:"):
                    continue
                body = entry[len("TODO:"):].strip()
                head, _, hint = body.partition(" - ")
                items.append((head, hint))
            blocks.append({"type": "todos", "items": items})
        elif not blocks and not lead:
            lead = " ".join(block)
        else:
            blocks.append({"type": "para", "spans": runs(" ".join(block))})

    return title, lead, blocks


HTML_PAGE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>{title}</title>
<style>
  :root {{
    --bg: #0c0c0e; --surface: #1b1b1f; --text: #f4f4f2; --muted: #a0a0a4;
    --accent: #ff7a18; --link: #ffb347; --todo: #ff4d9d; --border: #2a2a30;
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0; padding: clamp(1.5rem, 5vw, 3.5rem) 1.25rem;
    background: var(--bg); color: var(--text); line-height: 1.7;
    font-family: system-ui, -apple-system, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    -webkit-font-smoothing: antialiased;
  }}
  main {{ max-width: 820px; margin: 0 auto; }}
  a {{ color: var(--link); }}
  img {{ display: block; max-width: 100%; }}
  code {{ padding: .1em .35em; border-radius: 4px; background: rgba(255,255,255,.06); font-size: .88em; }}
  header {{ margin-bottom: 0; }}
  header img {{ width: min(340px, 70%); margin-bottom: 1.25rem; }}
  h1 {{ margin: 0 0 .35rem; font-size: clamp(1.5rem, 4vw, 2rem); letter-spacing: .04em; text-transform: uppercase; }}
  .lead {{ margin: 0; color: var(--muted); font-size: .95rem; }}
  h2 {{
    margin: 3rem 0 1rem; padding-top: 1.5rem; border-top: 1px solid var(--border);
    font-size: 1.15rem; letter-spacing: .1em; text-transform: uppercase;
  }}
  h3 {{ margin: 2rem 0 .6rem; font-size: .95rem; letter-spacing: .1em; text-transform: uppercase; }}
  h4 {{ margin: 1.6rem 0 .3rem; font-size: .85rem; letter-spacing: .14em; text-transform: uppercase; }}
  p {{ color: #d6d6d8; }}
  ul {{ padding-left: 1.15rem; color: #d6d6d8; }}
  li {{ margin-bottom: .4rem; }}
  li::marker {{ color: var(--accent); }}
  .facts {{
    display: grid; grid-template-columns: minmax(120px, 190px) 1fr; gap: 0;
    margin: 0 0 1.5rem; border: 1px solid var(--border); border-radius: 10px; overflow: hidden;
  }}
  .facts dt, .facts dd {{ margin: 0; padding: .6rem .9rem; border-top: 1px solid var(--border); }}
  .facts dt:first-of-type, .facts dt:first-of-type + dd {{ border-top: 0; }}
  .facts dt {{
    font-size: .7rem; font-weight: 700; letter-spacing: .12em; text-transform: uppercase;
    color: var(--muted); background: rgba(255,255,255,.02);
  }}
  .facts dd {{ font-size: .92rem; overflow-wrap: anywhere; }}
  blockquote {{
    margin: 0 0 1.5rem; padding: 1.1rem 1.35rem; background: var(--surface);
    border: 1px solid var(--border); border-left: 3px solid var(--accent); border-radius: 0 10px 10px 0;
  }}
  blockquote p {{ margin: 0; color: #e6e6e8; }}
  blockquote footer {{ margin-top: .6rem; color: var(--muted); font-size: .85rem; }}
  .files {{
    list-style: none; padding: 0; margin: 0 0 1.5rem;
    display: grid; grid-template-columns: repeat(auto-fill, minmax(260px, 1fr)); gap: .3rem .9rem;
  }}
  .files li {{ display: flex; justify-content: space-between; gap: 1rem; font-size: .85rem; margin: 0; }}
  .files span {{ flex: none; color: var(--muted); font-size: .78rem; }}
  .empty {{ color: var(--muted); font-size: .9rem; }}
  .soon {{ color: var(--todo); font-weight: 600; }}
  .clips {{
    display: grid; grid-template-columns: repeat(auto-fill, minmax(300px, 1fr));
    gap: 1rem; margin: 0 0 1.5rem;
  }}
  .clip {{ margin: 0; border: 1px solid var(--border); border-radius: 10px; background: var(--surface); overflow: hidden; }}
  .clip video {{ display: block; width: 100%; height: auto; background: #0a0a0c; }}
  .clip-meta {{ display: flex; flex-wrap: wrap; gap: .5rem; padding: .6rem .85rem; border-top: 1px solid var(--border); }}
  .clip-dl {{
    display: inline-flex; align-items: baseline; gap: .4rem; padding: .25rem .65rem;
    border: 1px solid var(--border); border-radius: 999px; color: var(--muted);
    text-decoration: none; font-size: .72rem; font-weight: 700; letter-spacing: .08em;
    text-transform: uppercase;
  }}
  .clip-dl span {{ font-weight: 500; letter-spacing: .02em; text-transform: none; opacity: .75; }}
  .todos {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(240px, 1fr)); gap: .6rem; margin: 0 0 1.5rem; }}
  .todo {{ padding: .8rem 1rem; border: 1px dashed var(--todo); border-radius: 10px; background: rgba(255,77,157,.08); }}
  .todo-tag {{
    display: inline-block; margin-bottom: .3rem; padding: .05rem .5rem; border-radius: 4px;
    background: var(--todo); color: #2a0316; font-size: .62rem; font-weight: 800;
    letter-spacing: .16em; text-transform: uppercase;
  }}
  .todo p {{ margin: 0; }}
  .todo-title {{ font-weight: 700; color: #fff; }}
  .todo-hint {{ color: #b9a4af !important; font-size: .85rem; }}
  @media (max-width: 640px) {{
    .facts {{ grid-template-columns: 1fr; }}
    .facts dt {{ padding-bottom: 0; }}
    .facts dt + dd {{ border-top: 0; padding-top: .1rem; }}
  }}
</style>
</head>
<body>
<main>
<header>
  <img src="logos/beltfed-logo-color.png" alt="BELTFED" />
  <h1>{title}</h1>
  <p class="lead">{lead}</p>
</header>
{body}
</main>
</body>
</html>
"""


def html_spans(spans):
    out = []
    for span in spans:
        text = html.escape(span["text"], quote=False)
        if span.get("code"):
            text = f"<code>{text}</code>"
        if span.get("bold"):
            text = f"<strong>{text}</strong>"
        if span.get("href"):
            text = f'<a href="{html.escape(span["href"], quote=True)}">{text}</a>'
        out.append(text)
    return "".join(out)


def render_html(title, lead, blocks, target):
    out = []
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            tag = f"h{block['level']}"
            out.append(f"<{tag}>{html_spans(block['spans'])}</{tag}>")
        elif kind == "para":
            out.append(f"<p>{html_spans(block['spans'])}</p>")
        elif kind == "rule":
            out.append("<hr />")
        elif kind == "list":
            out.append(
                "<ul>"
                + "".join(f"<li>{html_spans(i)}</li>" for i in block["items"])
                + "</ul>"
            )
        elif kind == "facts":
            rows = ""
            for key, value in block["rows"]:
                cell = (
                    f'<span class="soon">{html.escape(plain(value), quote=False)}</span>'
                    if is_soon(value)
                    else html_spans(value)
                )
                rows += f"<dt>{html.escape(key, quote=False)}</dt><dd>{cell}</dd>"
            out.append(f'<dl class="facts">{rows}</dl>')
        elif kind == "quote":
            body = f"<p>{html_spans(block['spans'])}</p>"
            if block["attribution"]:
                body += f"<footer>- {html.escape(block['attribution'], quote=False)}</footer>"
            out.append(f"<blockquote>{body}</blockquote>")
        elif kind == "todos":
            cards = ""
            for head, hint in block["items"]:
                cards += '<div class="todo"><span class="todo-tag">TODO</span>'
                cards += f'<p class="todo-title">{html_spans(runs(head))}</p>'
                if hint:
                    cards += f'<p class="todo-hint">{html_spans(runs(hint))}</p>'
                cards += "</div>"
            out.append(f'<div class="todos">{cards}</div>')
        elif kind == "clips":
            cards = ""
            for clip in block["clips"]:
                links = "".join(
                    f'<a class="clip-dl" href="{SITE_CLIPS}{name}">{label}'
                    f"<span>{format_bytes(size)}</span></a>"
                    for label, name, size in clip["files"]
                )
                cards += (
                    '<figure class="clip">'
                    f'<video src="{SITE_CLIPS}{clip["stem"]}.webm" autoplay loop muted '
                    'playsinline preload="metadata"></video>'
                    f'<figcaption class="clip-meta">{links}</figcaption>'
                    "</figure>"
                )
            out.append(f'<div class="clips">{cards}</div>')
        elif kind == "files":
            if not block["files"]:
                out.append(
                    f'<p class="empty">Nothing in <code>{block["folder"]}/</code> yet.</p>'
                )
            else:
                items = "".join(
                    f'<li><a href="{block["folder"]}/{name}">{name}</a>'
                    f"<span>{format_bytes(size)}</span></li>"
                    for name, size in block["files"]
                )
                out.append(f'<ul class="files">{items}</ul>')

    target.write_text(
        HTML_PAGE.format(
            title=html.escape(title, quote=False),
            lead=html.escape(lead, quote=False),
            body="".join(out),
        ),
        encoding="utf-8",
    )


def logo_on_white(max_width=1200):
    with Image.open(logo) as image:
        image = image.convert("RGBA")
        if image.width > max_width:
            height = round(image.height * max_width / image.width)
            image = image.resize((max_width, height), Image.LANCZOS)
        flat = Image.new("RGB", image.size, "white")
        flat.paste(image, mask=image.getchannel("A"))
        buffer = io.BytesIO()
        flat.save(buffer, "PNG", optimize=True)
        buffer.seek(0)
        return buffer, flat.width, flat.height


def pdf_markup(spans):
    out = []
    for span in spans:
        text = html.escape(span["text"], quote=False)
        if span.get("code"):
            text = f'<font face="Courier">{text}</font>'
        if span.get("bold"):
            text = f"<b>{text}</b>"
        if span.get("href"):
            href = html.escape(span["href"], quote=True)
            text = f'<a href="{href}" color="{ACCENT}"><u>{text}</u></a>'
        out.append(text)
    return "".join(out)


def pdf_styles():
    base = ParagraphStyle(
        "body", fontName="Helvetica", fontSize=9.5, leading=14,
        textColor=colors.HexColor(INK), spaceAfter=7,
    )
    return {
        "body": base,
        "title": ParagraphStyle("title", parent=base, fontName="Helvetica-Bold",
                                fontSize=17, leading=21, spaceAfter=2),
        "lead": ParagraphStyle("lead", parent=base, fontSize=9, spaceAfter=14,
                               textColor=colors.HexColor(MUTED)),
        2: ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=12.5,
                          leading=16, spaceBefore=16, spaceAfter=8,
                          textColor=colors.HexColor(ACCENT)),
        3: ParagraphStyle("h3", parent=base, fontName="Helvetica-Bold", fontSize=10,
                          leading=13, spaceBefore=11, spaceAfter=4),
        4: ParagraphStyle("h4", parent=base, fontName="Helvetica-Bold", fontSize=9,
                          leading=12, spaceBefore=9, spaceAfter=2),
        "small": ParagraphStyle("small", parent=base, fontSize=8, leading=11.5,
                                spaceAfter=9, textColor=colors.HexColor(MUTED)),
        "quote": ParagraphStyle("quote", parent=base, fontSize=10, leading=15,
                                leftIndent=9, spaceAfter=3),
        "attrib": ParagraphStyle("attrib", parent=base, fontSize=8.5, leftIndent=9,
                                 spaceAfter=10, textColor=colors.HexColor(MUTED)),
        "key": ParagraphStyle("key", parent=base, fontName="Helvetica-Bold", fontSize=8,
                              leading=11, spaceAfter=0, textColor=colors.HexColor(MUTED)),
        "value": ParagraphStyle("value", parent=base, fontSize=9, leading=12, spaceAfter=0),
        "todo": ParagraphStyle("todo", parent=base, fontSize=9, leading=12.5,
                               spaceAfter=0, textColor=colors.HexColor(TODO_INK)),
    }


def render_pdf(title, lead, blocks, target):
    style = pdf_styles()
    story = []

    image, width, height = logo_on_white()
    display = 58 * mm
    story.append(PdfImage(image, width=display, height=display * height / width))
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph(html.escape(title, quote=False), style["title"]))
    story.append(Paragraph(html.escape(lead, quote=False), style["lead"]))

    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            story.append(Paragraph(pdf_markup(block["spans"]), style[block["level"]]))
        elif kind == "para":
            story.append(Paragraph(pdf_markup(block["spans"]), style["body"]))
        elif kind == "rule":
            story.append(Spacer(1, 4 * mm))
        elif kind == "list":
            for item in block["items"]:
                story.append(Paragraph("&bull;&nbsp;&nbsp;" + pdf_markup(item), style["body"]))
        elif kind == "facts":
            rows = [
                [Paragraph(html.escape(key, quote=False).upper(), style["key"]),
                 Paragraph(
                     f'<font color="{TODO_INK}">{html.escape(plain(value), quote=False)}</font>'
                     if is_soon(value)
                     else pdf_markup(value),
                     style["value"],
                 )]
                for key, value in block["rows"]
            ]
            table = Table(rows, colWidths=[38 * mm, 122 * mm], hAlign="LEFT")
            table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(RULE)),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#FAF8F6")),
            ]))
            story.append(table)
            story.append(Spacer(1, 5 * mm))
        elif kind == "quote":
            group = [Paragraph(pdf_markup(block["spans"]), style["quote"])]
            if block["attribution"]:
                group.append(Paragraph(
                    "- " + html.escape(block["attribution"], quote=False), style["attrib"]
                ))
            story.append(KeepTogether(group))
        elif kind == "todos":
            rows = []
            for head, hint in block["items"]:
                text = f"<b>TODO</b>&nbsp;&nbsp;{pdf_markup(runs(head))}"
                if hint:
                    text += f" &ndash; {pdf_markup(runs(hint))}"
                rows.append([Paragraph(text, style["todo"])])
            table = Table(rows, colWidths=[160 * mm], hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(TODO_BG)),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(TODO_INK)),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#F3C6D9")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(table)
            story.append(Spacer(1, 5 * mm))
        elif kind == "clips":
            for clip in block["clips"]:
                links = "  ".join(
                    f'<a href="{SITE_CLIPS}{name}" color="{ACCENT}"><u>{label}</u></a> '
                    f"({format_bytes(size)})"
                    for label, name, size in clip["files"]
                )
                story.append(Paragraph(
                    f"<b>{html.escape(clip['stem'], quote=False)}</b> &ndash; {links}",
                    style["small"],
                ))
        elif kind == "files":
            if not block["files"]:
                story.append(Paragraph(f"Nothing in {block['folder']}/ yet.", style["small"]))
            else:
                names = html.escape(", ".join(n for n, _ in block["files"]), quote=False)
                total = format_bytes(sum(s for _, s in block["files"]))
                story.append(Paragraph(
                    f"<b>{len(block['files'])} files</b> in {block['folder']}/ "
                    f"({total}) &ndash; {names}",
                    style["small"],
                ))

    def page_number(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(canvas.getPageNumber()))
        canvas.restoreState()

    SimpleDocTemplate(
        str(target), pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title=title, author="Hello World Studios",
    ).build(story, onFirstPage=page_number, onLaterPages=page_number)


def docx_hyperlink(paragraph, url, text):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT.lstrip("#"))
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def docx_spans(paragraph, spans, size=None, color=None, italic=False):
    for span in spans:
        if span.get("href"):
            docx_hyperlink(paragraph, span["href"], span["text"])
            continue
        run = paragraph.add_run(span["text"])
        run.bold = bool(span.get("bold"))
        run.italic = italic
        if span.get("code"):
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def render_docx(title, lead, blocks, target):
    document = docx.Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    image, _, _ = logo_on_white()
    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.add_run().add_picture(image, width=Inches(2.3))

    heading = document.add_paragraph().add_run(title)
    heading.bold = True
    heading.font.size = Pt(19)
    heading.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))

    tagline = document.add_paragraph().add_run(lead)
    tagline.font.size = Pt(9)
    tagline.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))

    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            paragraph = document.add_heading(level=min(level, 4))
            docx_spans(
                paragraph, block["spans"],
                size={2: 14, 3: 11, 4: 10}.get(level, 10),
                color=ACCENT if level == 2 else INK,
            )
            for run in paragraph.runs:
                run.bold = True
        elif kind == "para":
            docx_spans(document.add_paragraph(), block["spans"])
        elif kind == "rule":
            document.add_paragraph()
        elif kind == "list":
            for item in block["items"]:
                docx_spans(document.add_paragraph(style="List Bullet"), item)
        elif kind == "facts":
            table = document.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for key, value in block["rows"]:
                cells = table.add_row().cells
                label = cells[0].paragraphs[0].add_run(key.upper())
                label.bold = True
                label.font.size = Pt(8)
                label.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
                if is_soon(value):
                    docx_spans(cells[1].paragraphs[0], [{"text": plain(value)}],
                               size=9.5, color=TODO_INK)
                else:
                    docx_spans(cells[1].paragraphs[0], value, size=9.5)
            document.add_paragraph()
        elif kind == "quote":
            paragraph = document.add_paragraph(style="Intense Quote")
            docx_spans(paragraph, block["spans"])
            if block["attribution"]:
                attribution = document.add_paragraph().add_run("- " + block["attribution"])
                attribution.italic = True
                attribution.font.size = Pt(9)
                attribution.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
        elif kind == "todos":
            for head, hint in block["items"]:
                paragraph = document.add_paragraph()
                tag = paragraph.add_run("TODO  ")
                tag.bold = True
                tag.font.color.rgb = RGBColor.from_string(TODO_INK.lstrip("#"))
                docx_spans(paragraph, runs(head), color=TODO_INK)
                if hint:
                    docx_spans(paragraph, [{"text": " - " + hint}],
                               color=TODO_INK, italic=True)
        elif kind == "clips":
            for clip in block["clips"]:
                paragraph = document.add_paragraph()
                docx_spans(paragraph, [{"text": clip["stem"] + "  ", "bold": True}],
                           size=8.5)
                for index, (label, name, size) in enumerate(clip["files"]):
                    if index:
                        docx_spans(paragraph, [{"text": "  "}], size=8.5, color=MUTED)
                    docx_spans(paragraph, [{"text": label, "href": SITE_CLIPS + name}])
                    docx_spans(paragraph, [{"text": f" ({format_bytes(size)})"}],
                               size=8.5, color=MUTED)
        elif kind == "files":
            paragraph = document.add_paragraph()
            if not block["files"]:
                docx_spans(paragraph, [{"text": f"Nothing in {block['folder']}/ yet."}],
                           size=8.5, color=MUTED, italic=True)
            else:
                total = format_bytes(sum(s for _, s in block["files"]))
                names = ", ".join(n for n, _ in block["files"])
                docx_spans(paragraph, [
                    {"text": f"{len(block['files'])} files", "bold": True},
                    {"text": f" in {block['folder']}/ ({total}) - {names}"},
                ], size=8.5, color=MUTED)

    document.save(str(target))


def collect(folder):
    directory = kit / folder
    if not directory.is_dir():
        return []
    return sorted(
        (f for f in directory.iterdir() if f.is_file() and f.suffix.lower() != ".txt"),
        key=lambda f: f.name,
    )


def write_manifest(listing, clips):
    manifest.write_text(
        json.dumps(
            {
                "folders": {
                    folder: [
                        {"name": f.name, "size": f.stat().st_size} for f in files
                    ]
                    for folder, files in listing.items()
                },
                "clips": [
                    {
                        "stem": clip["stem"],
                        "files": [
                            {"label": label, "name": name, "size": size}
                            for label, name, size in clip["files"]
                        ],
                    }
                    for clip in clips
                ],
            },
            indent=1,
        ),
        encoding="utf-8",
    )
    total = sum(len(v) for v in listing.values()) + len(clips)
    print(f"{manifest.name}: {total} entries")


def main():
    listing = {folder: collect(folder) for folder in ASSET_FOLDERS}
    write_manifest(listing, collect_clips())
    title, lead, blocks = parse(source.read_text(encoding="utf-8"), listing)

    outputs = [
        ("presskit.html", render_html, True),
        ("presskit.pdf", render_pdf, HAVE_PDF),
        ("presskit.docx", render_docx, HAVE_DOCX),
    ]
    for name, renderer, available in outputs:
        target = kit / name
        if not available:
            print(f"{name}: skipped, pip install reportlab python-docx")
            continue
        renderer(title, lead, blocks, target)
        print(f"{name}: {target.stat().st_size / 1024:.0f} KB from {source.name}")

    files = [
        path
        for path in sorted(kit.rglob("*"))
        if path.is_file() and path.name.lower() not in SKIP_NAMES
    ]
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for path in files:
            zf.write(path, str(Path(kit.name) / path.relative_to(kit)).replace("\\", "/"))
    print(f"{archive.name}: {len(files)} files, {archive.stat().st_size / 1048576:.1f} MB")


if __name__ == "__main__":
    main()
